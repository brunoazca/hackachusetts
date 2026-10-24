#!/usr/bin/env python3
"""Build and validate the offline Civic knowledge-base index from saved evidence.

No network. Optional pypdf/openpyxl enrich file previews when available. All input
is treated as data; downloaded scripts, formulas, macros and HTML are not run.
"""
import collections
import csv
import hashlib
import html
import io
import json
import os
import pathlib
import re
import sys
import urllib.parse
import zipfile
import xml.etree.ElementTree as ET

sys.dont_write_bytecode = True
from inventory_civic_portal import ROOT, OUT, SMALL, now, digest, profile_csv, save_json


def read(name):
    return json.loads((OUT / name).read_text())


def esc(value):
    return str(value if value is not None else '').replace('|', '\\|').replace('\n', ' ').replace('\r', ' ')


def local_link(path, label='Local file', origin=None):
    origin = origin or OUT
    target = os.path.relpath(ROOT / path, origin)
    return f'[{esc(label)}](<{target}>)'


def source_link(url, label='Source'):
    return f'[{esc(label)}](<{url}>)'


def load_requests():
    rows = []
    for name in ['requests', 'followup_requests', 'nested_requests', 'ckan_requests',
                 'crash_requests', 'additional_requests', 'layer_requests', 'daily_bike_requests']:
        path = OUT / (name + '.json')
        if path.exists():
            rows.extend(json.loads(path.read_text()))
    for item in read('socrata_inventory.json'):
        rows.extend(item['requests'])
    # Retain original observations, but normalize curl's equivalent size-cap errors.
    for row in rows:
        if re.search(r'(file size exceeded|maximum allowed file size)',row.get('error') or '',re.I):
            row['access'] = 'large_remote_file'
            row['retention'] = 'preview_or_metadata'
        if row.get('local_path'):
            path = ROOT / row['local_path']
            if not path.exists():
                raise ValueError('Missing response file: '+str(path))
            if digest(path.read_bytes()) != row['local_sha256']:
                raise ValueError('Response hash mismatch: '+str(path))
    return rows


def inspect_files(requests):
    try:
        from pypdf import PdfReader
    except ImportError:
        PdfReader = None
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    by_path = {r['local_path']:r for r in requests if r.get('local_path')}
    files = []
    for path in sorted((OUT / 'downloads').glob('*')):
        rel = str(path.relative_to(ROOT))
        request = by_path.get(rel,{})
        entry = {'path':rel,'name':request.get('name',path.stem),'bytes':path.stat().st_size,
                 'sha256':digest(path.read_bytes()),'source_url':request.get('url')}
        try:
            if path.suffix == '.pdf' and PdfReader:
                pdf = PdfReader(path)
                entry['pages'] = len(pdf.pages)
                entry['preview'] = ' '.join((pdf.pages[0].extract_text() or '').split()[:25]) if pdf.pages else ''
                entry['pdf_title'] = str((pdf.metadata or {}).get('/Title',''))
            elif path.suffix == '.xlsx' and openpyxl:
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                entry['sheets'] = [{'title':sheet.title,'grid_rows':sheet.max_row,'grid_columns':sheet.max_column,
                    'preview':[list(row) for row in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row or 1,6),
                        max_col=min(sheet.max_column or 1,40),values_only=True)]} for sheet in book]
                entry['note'] = 'Grid dimensions may include formatting and blanks. Values are cached source values; formulas were not evaluated.'
                book.close()
            elif path.suffix == '.docx':
                with zipfile.ZipFile(path) as z:
                    xml = ET.fromstring(z.read('word/document.xml'))
                    texts = [n.text or '' for n in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')]
                    entry['preview'] = ' '.join(' '.join(texts).split()[:25])
                    entry['tables'] = len(list(xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl')))
            elif path.suffix == '.csv':
                profile = profile_csv(path)
                entry.update({k:v for k,v in profile.items() if k != 'sample'})
                entry['preview'] = profile['sample'][:5]
            elif path.suffix == '.zip':
                with zipfile.ZipFile(path) as z:
                    entry['members'] = [{'path':f.filename,'bytes':f.file_size} for f in z.infolist()]
                    # Full archive is retained; inspect CSV schemas without extracting paths.
                    csv_members=[]
                    for member in z.infolist():
                        if member.filename.lower().endswith('.csv') and member.file_size <= 50*1024*1024:
                            with z.open(member) as f:
                                reader=csv.reader(io.TextIOWrapper(f,encoding='utf-8-sig',errors='replace'))
                                columns=next(reader,[]); count=sum(1 for _ in reader)
                            csv_members.append({'path':member.filename,'columns':columns,'rows':count})
                    entry['csv_tables']=csv_members
        except Exception as e:
            entry['inspection_error'] = type(e).__name__+': '+str(e)
        files.append(entry)
    # JSON encoder default handles spreadsheet datetime values without mutating originals.
    (OUT/'file_inventory.json').write_text(json.dumps(files,ensure_ascii=False,indent=2,default=str)+'\n')
    return files


def validate_kits(portal, requests):
    req = {r['id']:r for r in requests}
    checks, profiles = [], []
    with zipfile.ZipFile(ROOT/'data/civic/kits/tracks-civic.zip') as track:
        members={p.split('/')[-2]:p for p in track.namelist() if p.endswith('/data.csv')}
        assert set(members)==set(portal['track_kit']['datasets']), 'Track membership mismatch'
        for key in portal['track_kit']['datasets']:
            path=ROOT/'data/civic/datasets'/key/'data.csv'
            p=profile_csv(path)
            p['id']=key
            p['path']=str(path.relative_to(ROOT))
            sample=p.pop('sample')
            sample_path=OUT/'previews'/(key+'.json')
            save_json(sample_path,{'dataset':key,'source':p['path'],'source_sha256':p['sha256'],
                'total_snapshot_rows':p['rows'],'preview_rows':len(sample),'selection':'First 100 file rows; not a representative sample',
                'columns':p['columns'],'records':sample})
            p['preview_path']=str(sample_path.relative_to(ROOT))
            assert digest(track.read(members[key]))==p['sha256'], key+' differs from original track ZIP'
            with zipfile.ZipFile(ROOT/req['kit-'+key]['local_path']) as kit:
                csv_members=[m for m in kit.namelist() if m.endswith('/data.csv') or m=='data.csv']
                assert len(csv_members)==1, key+' dataset kit has unexpected CSV members'
                assert digest(kit.read(csv_members[0]))==p['sha256'], key+' individual kit differs'
            profiles.append(p)
        for key,kit_info in portal['challenge_kits'].items():
            path=ROOT/req['challenge-'+key]['local_path']
            with zipfile.ZipFile(path) as z:
                sub={p.split('/')[-2]:p for p in z.namelist() if p.endswith('/data.csv')}
                assert set(sub)==set(kit_info['datasets']),key+' challenge kit membership differs'
                for name,path_in_zip in sub.items():
                    assert digest(z.read(path_in_zip))==digest(track.read(members[name])),key+'/'+name+' contents differ'
            checks.append({'challenge':key,'datasets':list(sub),'all_csv_hashes_match_track':True,
                           'path':req['challenge-'+key]['local_path']})
        assert set().union(*(set(c['datasets']) for c in checks))==set(members),'Challenge union mismatch'
    save_json(OUT/'starter_profiles.json',profiles)
    return profiles,checks


def write_resource_pages(portal, profiles, requests):
    req={r['id']:r for r in requests}; prof={p['id']:p for p in profiles}
    directory=OUT/'resources';directory.mkdir(exist_ok=True)
    for r in portal['resources']:
        response=req[r['id']]
        lines=[f"# {r['name']}",'',r['summary'],'',f"Use: {r['usefulFor']}",'',
               '- '+source_link(r['href']),f"- Kind: {r['kind']}",
               '- Portal challenge tags: '+(', '.join(c.upper() for c in r['challenges']) or 'Shared / unassigned'),
               f"- Source request: HTTP {response['http_status']}; {response['access']}; {response['checked_at_utc']}",
               '- Retention: '+response.get('retention','access result only')]
        if r.get('kitId'):
            p=prof[r['kitId']]
            memberships=[c.upper() for c,v in portal['challenge_kits'].items() if r['kitId'] in v['datasets']]
            lines += ['','## Local starter snapshot','',f"{p['rows']:,} rows; {len(p['columns'])} columns; {p['bytes']:,} bytes. Complete organizer file, not a claim of a complete upstream database.",'',
                '- '+local_link(p['path'],'Full starter CSV',directory),
                '- '+local_link(p['preview_path'],'100-row preview',directory),
                '- '+local_link(f"data/civic/datasets/{r['kitId']}/README.md",'Organizer README and query',directory),
                '- Kit membership: '+', '.join(memberships),
                '- SHA-256: `'+p['sha256']+'`','','### Columns','',
                '| Column | Blank cells in snapshot |','|---|---:|']
            lines += [f"| `{esc(c)}` | {p['blank_cells'].get(c,0):,} |" for c in p['columns']]
        lines += ['','## Source response','']
        if response.get('local_path'):
            lines.append(local_link(response['local_path'],'Read captured source response',directory))
        elif response.get('page_title'):
            lines += ['Page title: '+response['page_title'], '', 'Short opening excerpt (may include navigation): '+response.get('excerpt','')]
        else:
            lines.append(response.get('error') or 'The HTTP response did not provide a retained data file. See the access ledger.')
        if response.get('returned_rows') is not None:
            lines.append(f"Returned rows: {response['returned_rows']:,}. API limits and filters are recorded in the exact source URL.")
        if response.get('reported_total') is not None:
            lines.append(f"API-reported total: {response['reported_total']:,} (at fetch time).")
        related=[x for x in requests if x.get('parent')==r['id']]
        if related:
            lines += ['','## Inspected subsets and related requests','','| Resource | Result | Local evidence |','|---|---|---|']
            for x in related:
                local=local_link(x['local_path'],'File',directory) if x.get('local_path') else 'Preview / status in manifest'
                lines.append(f"| {source_link(x['url'],x['name'])} | {x['access']} | {local} |")
        if r['id']=='e2c':
            lines += ['','The command-line 403 was resolved by reading the directory in the browser. All 133 linked asset IDs are inventoried in [E2C datasets and dashboards](../E2C.md).']
        if r['id'] in ('crash-background','civic-extra'):
            lines += ['','The Google Drive connector successfully read this document. Its embedded hyperlinks were extracted separately; full third-party prose and embedded images are not republished. See [supplementary files](../SUPPLEMENTARY.md).']
        lines += ['','## Limits','',
            'A successful GET verifies a response, not a working application integration. First-page previews are not representative samples. Counts, units, dates, license terms and missing values need to be interpreted for the chosen product.',
            '', 'All discovered page links, including navigation, are preserved in `requests` inside [manifest.json](../manifest.json). They are not all recursively downloaded.','']
        (directory/(r['id']+'.md')).write_text('\n'.join(lines))


def write_e2c(items):
    lines=['# E2C: every dataset and dashboard linked from the directory','',
        'Directory read in the browser on October 24, 2026. Public API schemas, reported row counts and response timestamps are in [socrata_inventory.json](socrata_inventory.json). Grid dimensions, source row counts and saved sample rows are distinct measures.','',
        'A full table means the CSV parsed row count equals the separately queried source count. Tables up to 50,000 rows are attempted in full, with a 5 MiB download limit; larger tables may also be retained when estimated small. Other tables have up to 100 source rows. Stories and dashboards have descriptive metadata and a landing-page preview, not an asserted export.','',
        '| Resource | Source rows | Retention | Data or preview | Schema |','|---|---:|---|---|---|']
    for item in items:
        data=item.get('data_path') or item.get('preview_path')
        local=local_link(data,'Open') if data else 'Metadata / landing page'
        schema=next((r.get('local_path') for r in item['requests'] if r['role']=='socrata_schema'),None)
        lines.append(f"| {source_link(item['page_url'],item.get('name',item['id']))} | {item.get('source_rows','—')} | {item['retention']} | {local} | {local_link(schema,'Fields') if schema else item['access']} |")
    lines += ['','The public catalog contains additional assets beyond the directory. [Catalog metadata](responses/nested-6a33f3ce33de.json) inventories all 232 entries returned by that API; those extra entries are discovery metadata, not 232 locally copied datasets.','']
    (OUT/'E2C.md').write_text('\n'.join(lines))


def write_families(requests):
    req={r['id']:r for r in requests};lines=['# Boston source database families and subsets','',
        'All resources listed by the six CKAN packages linked from the Civic starter datasets. Large raw files are represented by their exact download URL, source metadata and a bounded API preview.','']
    for family in read('ckan_families.json'):
        lines += ['## '+family['title'],'',
            'License reported by publisher: '+str(family.get('license_title') or 'unspecified')+'. '+(family.get('license_url') or ''),'',
            '| Resource | Format | Raw download | API preview | Source total |','|---|---|---|---|---:|']
        for f in family['resources']:
            raw=req.get('ckan-'+f['id'],{});preview=req.get('ckan-preview-'+f['id'],{})
            name=f.get('name') or urllib.parse.unquote(urllib.parse.urlparse(f['url']).path.split('/')[-1])
            local=local_link(raw['local_path'],'Full file') if raw.get('local_path') else raw.get('access','not downloaded')
            sample=local_link(preview['local_path'],str(preview.get('returned_rows','?'))+' rows') if preview.get('local_path') else 'Not a datastore table'
            total=preview.get('reported_total','—')
            lines.append(f"| {source_link(f['url'],name)} | {esc(f.get('format'))} | {local} | {sample} | {total} |")
        lines.append('')
    lines += ['## Empty datastore observations','',
        'The 2011 and 2014 Boston 311 datastore endpoints returned zero records while their raw CSV downloads are advertised as large files. This is an access/data-store discrepancy, not proof that no cases occurred. Raw file URLs and metadata remain recorded.','']
    (OUT/'BOSTON_FAMILIES.md').write_text('\n'.join(lines))


def write_supplementary(files,requests,sheets):
    lines=['# Reports, supplementary references, and spreadsheet tabs','',
        'Full small report/workbook files are retained as downloaded. Page counts, worksheet names and bounded opening previews are in [file_inventory.json](file_inventory.json). Binary files were inspected without running embedded code or formulas.','',
        '## Downloaded files','','| File / subject | Contents | Bytes |','|---|---|---:|']
    for f in files:
        if f['path'].endswith('.csv') or '/kit-' in f['path'] or '/challenge-' in f['path'] or '/track-' in f['path']:
            continue
        content=str(f['pages'])+' PDF pages' if 'pages' in f else ', '.join(s['title'] for s in f.get('sheets',[])) or (str(f['tables'])+' Word tables' if 'tables' in f else pathlib.Path(f['path']).suffix)
        lines.append(f"| {local_link(f['path'],f['name'])} | {esc(content)} | {f['bytes']:,} |")
    lines += ['','## Crash analysis Google Sheets','','| Workbook / tab | Captured range | Retention | Values |','|---|---|---|---|']
    for s in sheets:
        lines.append(f"| {source_link(s['url'],s['sheet'])} | {s['range']} | {s['retention']} | {local_link(s['local_path'],'Open')} |")
    lines += ['','Google Sheets captures contain formatted values; formulas, charts, formatting and image content are not reconstructed. Small tabs were read through their full allocated grid. Larger tabs have bounded rows with all columns.','',
        '## Supplementary document links','',
        'The portal’s additional-resource document is titled *Ideas and resources for the 9/28/2025 hackathon*. It is historical accessibility guidance, not a 2026 dataset release. The crash background is a collaborator’s analysis and method notes, not an official reconciliation of crash systems.','',
        '| Reference | Access | Local file |','|---|---|---|']
    for r in requests:
        if r['role'] in ('crash_background_link','supplementary_document_link','linked_report_or_subset','e2c_directory_link'):
            local=local_link(r['local_path'],'Open') if r.get('local_path') else 'Preview / status in manifest'
            lines.append(f"| {source_link(r['url'],r['name'])} | {r['access']} | {local} |")
    (OUT/'SUPPLEMENTARY.md').write_text('\n'.join(lines)+'\n')


def write_search_html(portal,profiles,requests,items):
    req={r['id']:r for r in requests};profile={p['id']:p for p in profiles}
    rows=[]
    for r in portal['resources']:
        p=profile.get(r.get('kitId'),{})
        rows.append({'name':r['name'],'group':'Civic portal','kind':r['kind'],
            'description':r['summary'],'status':'full starter snapshot' if p else req[r['id']]['access'],
            'rows':p.get('rows'),'link':'resources/'+r['id']+'.md','source':r['href'],
            'tags':', '.join(c.upper() for c in r['challenges']), 'fields':' '.join(p.get('columns',[]))})
    for r in items:
        path=r.get('data_path') or r.get('preview_path')
        rows.append({'name':r.get('name',r['id']),'group':'E2C','kind':r.get('view_type','reference'),
            'description':re.sub('<[^>]+>','',r.get('description') or '')[:350],
            'status':r['retention'],'rows':r.get('source_rows'),
            'link':os.path.relpath(ROOT/path,OUT) if path else 'E2C.md','source':r['page_url'],'tags':'',
            'fields':' '.join(str(c.get('field',''))+' '+str(c.get('name','')) for c in r.get('columns',[]))})
    for r in requests:
        if r['role'] in ('portal_resource','dataset_kit','socrata_schema','socrata_count','socrata_preview','socrata_full_csv','socrata_story_or_dashboard'):
            continue
        rows.append({'name':r['name'],'group':r.get('parent') or r['role'],'kind':r['role'],
            'description':r.get('page_title') or '', 'status':r['access'],'rows':r.get('reported_total',r.get('returned_rows')),
            'link':os.path.relpath(ROOT/r['local_path'],OUT) if r.get('local_path') else 'manifest.json','source':r['url'],'tags':''})
    # Escape script terminators in untrusted remote text, and render with textContent.
    payload=json.dumps(rows,ensure_ascii=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    template='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Civic source inventory</title>
<style>body{margin:0;font:16px/1.5 system-ui,sans-serif;background:#f4f6f8;color:#17212b}main{max-width:1180px;margin:auto;padding:40px 24px}h1{font-size:36px;margin:8px 0}p{max-width:900px}nav{display:flex;gap:18px;flex-wrap:wrap}a{color:#135bb0}label{display:block;margin-top:24px;font-weight:600}input{box-sizing:border-box;width:100%;padding:14px;border:1px solid #adb9c6;border-radius:8px;font:inherit}#count{color:#526375;margin:16px 0}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(310px,1fr));gap:16px}article{background:white;border:1px solid #d6dfe8;border-radius:10px;padding:20px;overflow-wrap:anywhere}h2{font-size:18px;margin:8px 0}.meta{font-size:13px;color:#526375}.status{font-weight:600}article p{font-size:14px}article a{display:inline-block;margin-right:18px}a:focus,input:focus{outline:3px solid #e5a315;outline-offset:3px}</style>
<main><div class="meta">HACKACHUSETTS · OCTOBER 24, 2026</div><h1>Civic source inventory</h1><p>Search datasets, challenge subsets, reports, and source responses. Full starter snapshots and small source tables are local. Larger data has bounded previews. Access errors and dashboard-only sources are labeled.</p>
<nav><a href="README.md">Inventory guide</a><a href="E2C.md">E2C</a><a href="BOSTON_FAMILIES.md">Boston source families</a><a href="SUPPLEMENTARY.md">Reports & tabs</a><a href="manifest.json">Machine-readable manifest</a></nav>
<label for="search">Find a resource, field of interest, source family, or challenge</label><input id="search" type="search" placeholder="Try 311, tuition, V3, crash, Numina, or full_dataset"><div id="count" aria-live="polite"></div><div id="results" class="grid"></div></main>
<script>const rows=__DATA__;const search=document.getElementById('search'),results=document.getElementById('results'),count=document.getElementById('count');function draw(){const q=search.value.toLowerCase().trim();const selected=rows.filter(r=>Object.values(r).join(' ').toLowerCase().includes(q));count.textContent=selected.length+' of '+rows.length+' catalog entries';results.replaceChildren();for(const r of selected){const card=document.createElement('article');for(const [tag,cls,text] of [['div','meta',r.group+' · '+r.kind+(r.tags?' · '+r.tags:'')],['h2','',r.name],['p','',r.description],['div','status',r.status+(r.rows!==null&&r.rows!==undefined?' · '+Number(r.rows).toLocaleString()+' rows':'')]]){const e=document.createElement(tag);e.className=cls;e.textContent=text;card.append(e)}for(const [text,url] of [['Local evidence',r.link],['Source',r.source]]){if(!url||(!url.startsWith('https://')&&url.includes(':')))continue;const a=document.createElement('a');a.href=url;a.textContent=text;card.append(a)}results.append(card)}}search.addEventListener('input',draw);draw();</script></html>'''
    (OUT/'index.html').write_text(template.replace('__DATA__',payload))
    return len(rows)


def main():
    portal=read('portal.json');requests=load_requests();items=read('socrata_inventory.json')
    profiles,kit_checks=validate_kits(portal,requests)
    files=inspect_files(requests)
    # Normalize the successful retry for the sheet whose name has trailing space.
    sheets=[]
    for path in sorted((OUT/'responses').glob('crash-sheet-*.json')):
        s=json.loads(path.read_text());s['local_path']=str(path.relative_to(ROOT))
        if s['id']=='crash-sheet-0-1260855922':
            s=read('metadata/crash-sheet-notes-retry.json');s['local_path']='data/civic/inventory/metadata/crash-sheet-notes-retry.json'
        if s.get('result',{}).get('error'):
            s['retention']='unavailable'
        sheets.append(s)
    daily=[read('responses/boston-daily-bike-full-'+str(n)+'.json') for n in (0,1000,2000)]
    features=[f for page in daily for f in page.get('features',[])]
    expected=read('responses/layer-cf076aad2170-count.json')['count']
    assert len(features)==expected,'Daily bicycle count pagination mismatch'
    oid=daily[0].get('objectIdFieldName') or 'OBJECTID'
    assert len({f['attributes'][oid] for f in features})==expected,'Duplicate bicycle feature IDs'
    full_daily={**daily[0],'features':features,'exceededTransferLimit':False}
    save_json(OUT/'responses/boston-daily-bike-complete.json',full_daily)
    write_resource_pages(portal,profiles,requests)
    write_e2c(items)
    write_families(requests)
    write_supplementary(files,requests,sheets)
    search_count=write_search_html(portal,profiles,requests,items)
    stats={'portal_resources':len(portal['resources']),'starter_datasets':len(profiles),
        'starter_rows':sum(p['rows'] for p in profiles),'challenge_kits':len(kit_checks),
        'e2c_assets':len(items),'e2c_retention_counts':dict(collections.Counter(x['retention'] for x in items)),
        'boston_families':len(read('ckan_families.json')),'boston_family_resources':sum(len(f['resources']) for f in read('ckan_families.json')),
        'request_observations':len(requests),'unique_requested_urls':len({r['url'] for r in requests}),
        'http_access_counts':dict(collections.Counter(r['access'] for r in requests)),
        'retained_download_files':len(files),'google_sheet_tabs':len(sheets),'daily_bike_full_records':expected,
        'search_entries':search_count}
    manifest={'generated_at_utc':now(),'source':portal['source_url'],'scope':
        'All 74 Civic cards and all 28 kits; five challenge subsets; all 43 resources in six linked Boston CKAN families; '
        'all 133 IDs from the E2C directory; direct report/document hyperlinks and four discovered dashboard data layers. '
        'General navigation links are inventoried but not recursively crawled across entire government websites. '
        'The 232-entry E2C catalog and Boston package list are discovery metadata, not claims of local dataset coverage.',
        'policy':{'max_new_raw_file_bytes':SMALL,'preview_rows':100,'small_socrata_attempt_rows':50000,
            'small_sheet_grid_cells':50000,'existing_larger_starter_files':'preserved intact',
            'web_pages':'title, short opening excerpt, links and access results; no full site/ranking mirror',
            'refresh':'explicit rerun only; source timestamps differ from observation periods'},
        'statistics':stats,'starter_profiles':profiles,'kit_checks':kit_checks,'requests':requests,
        'socrata_assets':[{k:v for k,v in x.items() if k!='requests'} for x in items],
        'google_sheet_tabs':[{k:v for k,v in x.items() if k not in ('result','tail_check')} for x in sheets],
        'browser_fallback':read('metadata/e2c-browser-directory.json'),
        'document_fallbacks':[read('metadata/civic-extra-drive.json'),read('metadata/crash-background-drive.json')],
        'full_daily_bicycle_layer':{'path':'data/civic/inventory/responses/boston-daily-bike-complete.json','rows':expected,
            'sha256':digest((OUT/'responses/boston-daily-bike-complete.json').read_bytes()),'method':'Three ordered pages; row count and unique OBJECTID count equal source count.'}}
    save_json(OUT/'manifest.json',manifest)
    req={r['id']:r for r in requests}
    guide=['# Civic data and source inventory','',
        'Captured October 24, 2026 from the [Civic portal](https://hackachusetts2026.vercel.app/#/civic). This inventory identifies exactly what is saved locally, what is a preview, and what could not be retrieved.','',
        f"**{stats['portal_resources']} portal resources · {stats['starter_datasets']} full starter datasets · {stats['starter_rows']:,} starter rows · {stats['challenge_kits']} challenge kits.**",'',
        '- [Searchable offline inventory](index.html)',
        '- [All E2C datasets and dashboards](E2C.md)',
        '- [Boston database families and every listed subset](BOSTON_FAMILIES.md)',
        '- [Reports, supplementary links, and spreadsheet tabs](SUPPLEMENTARY.md)',
        '- [Machine-readable manifest](manifest.json)',
        '- [Downloaded file contents, pages, tabs and schemas](file_inventory.json)','',
        '## What “full” means','',
        'All 22 organizer CSVs were already in the repository and remain intact, including larger files. Every individual dataset kit was fetched and its CSV hash checked against the track kit and local CSV. Every challenge ZIP was checked for exact membership and matching CSV contents. A full starter snapshot is still only the organizer’s slice of the original database.','',
        f"Additional E2C data: **{stats['e2c_retention_counts'].get('full_dataset',0)} complete tables**, **{stats['e2c_retention_counts'].get('preview',0)} table previews**, and **37 dashboard/reference entries**. The directory has 133 unique IDs. One story’s schema request returned 403; its public landing page is recorded separately. Full CSV counts were matched against source count queries. All 43 Boston-family resources have download attempts and/or source previews. The daily Boston bicycle layer is complete at 2,934 unique records.",'',
        'New small files are retained up to 5 MiB each. Tabular E2C resources with at most 50,000 rows are attempted in full; large tables receive first-page samples. Previews are not random or representative. Native dashboards and licensed rankings retain source descriptions and access evidence instead of invented bulk extracts. See the manifest for exact retrieval times and filters.','',
        '## Challenge subsets','', '| Challenge | Complete kit | Datasets in ZIP |','|---|---|---|']
    for c in portal['challenges']:
        kit=portal['challenge_kits'][c['id']]
        guide.append(f"| {c['code']} — {c['title']} | {local_link(req['challenge-'+c['id']]['local_path'],'ZIP')} | {', '.join(kit['datasets'])} |")
    guide += ['','The website’s per-card challenge tags and the ZIP membership are recorded separately; they are not assumed to be identical.','',
              '## Every Civic resource','','| Resource | Kind | Challenge tags | Saved data / access |','|---|---|---|']
    for r in portal['resources']:
        p=next((p for p in profiles if p['id']==r.get('kitId')),None)
        state=f"Full starter: {p['rows']:,} rows" if p else req[r['id']]['access']+'; reference preview'
        if r['id']=='e2c':state='Browser verified; 133 assets inventoried'
        if r['id'] in ('civic-extra','crash-background'):state='Drive verified; all document links inventoried'
        guide.append(f"| [{r['name']}](resources/{r['id']}.md) | {r['kind']} | {', '.join(c.upper() for c in r['challenges']) or 'Shared'} | {state} |")
    guide += ['','## Coverage boundaries and unresolved access','',
        'This is a complete inventory of the Civic portal’s 74 cards and five challenge bundles, plus the explicitly enumerated nested sources above. It is not a copy of every dataset or service linked somewhere inside Boston.gov, Mass.gov, Census, OpenStreetMap, Google Scholar, or every institution’s website. Discovered navigation links remain in the manifest. Private data, paid rankings, login-only analytics, and non-exported dashboards remain references.','',
        'Some Mass.gov, BLS, QS and other requests returned 403 or network errors. E2C and the two Google Docs were recovered through alternate read methods. Access results are observations from this environment, not proof that a source is gone. Historical 311 resources for 2011 and 2014 returned empty datastore responses despite large raw CSVs. Their raw URLs remain available. The linked DHE access/affordability URL returned 404; the E2C directory supplies current public-postsecondary alternatives.','',
        'The legacy BLS Massachusetts URL redirected to the broader OEWS tables page in the web research tool. Treat it as a navigation reference, not an archived Massachusetts wage table. Large Scorecard and Bluebikes originals are represented by their full organizer slices and 100-row local previews.','',
        '## Reproduce','',
        'From the repository root, use Python 3 and curl. Download the current page and its referenced JavaScript bundle, then parse the literal catalog without executing it:','',
        '```sh','python3 knowledge/sources/inventory_civic_portal.py --bundle /path/to/portal.js --initial',
        'python3 knowledge/sources/inventory_civic_portal.py --socrata data/civic/inventory/socrata_entries.json',
        'python3 knowledge/sources/inventory_civic_portal.py --entries data/civic/inventory/ckan_entries.json --output-name ckan_requests',
        'python3 knowledge/sources/build_civic_inventory.py','```','',
        'Other reviewed request lists end in `_entries.json`; pass the matching `--output-name` used by the builder. Requests use a temporary URL-keyed cache; delete only that task cache or call the helper with `refresh=True` for a new observation. Optional bundled `pypdf` and `openpyxl` add PDF page counts and worksheet previews. Browser and connected-Drive observations are separately dated evidence and are not silently regenerated. No downloaded code, macros, form submissions, or live database writes are executed.','']
    (OUT/'README.md').write_text('\n'.join(guide))
    knowledge=ROOT/'knowledge/CIVIC_INVENTORY.md'
    knowledge.write_text('# Civic knowledge base: complete source inventory\n\n'
        'Captured **October 24, 2026**. Start with the [inventory guide](../data/civic/inventory/README.md) or [searchable offline catalog](../data/civic/inventory/index.html).\n\n'
        f"- **74 Civic resources** cataloged with an access attempt for each, including all **22 full starter CSVs** ({stats['starter_rows']:,} rows).\n"
        '- **Five challenge kits** and **22 individual dataset kits**, with membership and CSV hashes checked against the full track kit.\n'
        '- **43 source resources** across six Boston database families, including annual 311 files, Numina sensor subsets, documentation, and neighborhood geometry.\n'
        f"- **133 E2C assets:** {stats['e2c_retention_counts'].get('full_dataset',0)} full tables, {stats['e2c_retention_counts'].get('preview',0)} data previews, and 37 dashboard/reference entries.\n"
        '- **Four dashboard data layers**, with schemas and previews; all **2,934 Boston daily bicycle-count records** are also saved.\n'
        '- Both supplementary Google Docs were read, their **40 distinct document links** enumerated (12 accessibility + 28 crash-background), and **11 tabs** across the linked crash-analysis spreadsheets captured fully or as previews.\n\n'
        'Every entry identifies its source, local path, access outcome, retrieval time, and whether the local content is full or sampled. The [manifest](../data/civic/inventory/manifest.json) and [file inventory](../data/civic/inventory/file_inventory.json) expose exact schemas, counts, hashes, PDF pages, and workbook tabs.\n\n'
        '**Limits:** “Full” refers to the named file or table, not an entire government website. New raw downloads are capped at 5 MiB. Large sources use bounded previews. Licensed, blocked, private, or non-exported sources are explicitly marked; their contents are not invented. General site navigation is indexed but not recursively mirrored. See the inventory guide for the access gaps and precise crawl boundary.\n\n'
        'Use [dataset verification](07_civic_dataset_verification.md) alongside this inventory before selecting joins, date ranges, or product claims. The source archive expands what we can inspect; it does not establish causal conclusions or production integrations.\n')
    # All retained response files and manually captured connector files get a checksum register.
    checksums={str(p.relative_to(ROOT)):digest(p.read_bytes()) for directory in ['responses','downloads','metadata','previews'] for p in (OUT/directory).glob('*') if p.is_file()}
    save_json(OUT/'checksums.json',checksums)
    assert len({r['id'] for r in portal['resources']})==74
    assert len(items)==len(read('metadata/e2c-browser-directory.json')['socrataIds'])
    for item in items:
        if item['retention']=='full_dataset':
            assert item['source_rows']==profile_csv(ROOT/item['data_path'])['rows']
    for path in OUT.rglob('*.json'):
        json.loads(path.read_text())
    save_json(OUT/'validation.json',{'checked_at_utc':now(),'passed':True,
        'checks':['74 unique portal resources','22 starter files match track and individual kits',
                  'five challenge kits have exact membership and matching CSV hashes','challenge union covers all starter datasets',
                  '133 E2C IDs inventoried','full E2C table counts equal observed API counts',
                  'daily bicycle layer count and unique IDs equal 2934','saved response SHA-256 checks','all inventory JSON parses'],
        'statistics':stats})
    print(json.dumps(stats,indent=2))


if __name__=='__main__':
    main()
